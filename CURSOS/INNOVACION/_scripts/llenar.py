"""Utilidades para llenar las guías de taller sin romper su formato."""
import docx, copy

EQUIPO = [
    "Cary Bernal, Tomás Gabriel  (coordinador)",
    "Zela Contreras, Heidi Nicole",
    "Rosas Quispe, Pernel Sebastián",
    "Apaza Cano, Daniela Grecia",
]
SEMESTRE = "2026-02"


def _set(cell, texto):
    """Escribe en una celda conservando el formato del primer run existente."""
    p = cell.paragraphs[0]
    if p.runs:
        p.runs[0].text = texto
        for r in p.runs[1:]:
            r.text = ""
    else:
        p.add_run(texto)
    for extra in cell.paragraphs[1:]:
        for r in extra.runs:
            r.text = ""


def set_cell(tabla, fila, col, texto):
    _set(tabla.rows[fila].cells[col], texto)


def cabecera(doc, semestre=SEMESTRE, grupo="", fecha="", hora=""):
    """Tabla 0: nombres del equipo y datos de entrega."""
    t = doc.tables[0]
    for i, nombre in enumerate(EQUIPO):
        if i < len(t.rows):
            _set(t.rows[i].cells[1], f"{i+1}.- {nombre}")
    for fila in t.rows:
        etq = fila.cells[0].text.strip().lower()
        if etq.startswith("semestre") and semestre:
            _set(fila.cells[1], semestre)
        elif etq.startswith("n° de grupo") and grupo:
            _set(fila.cells[1], grupo)
        elif etq.startswith("fecha de entrega") and fecha:
            _set(fila.cells[1], fecha)


def set_parrafo(doc, idx, texto):
    p = doc.paragraphs[idx]
    if p.runs:
        p.runs[0].text = texto
        for r in p.runs[1:]:
            r.text = ""
    else:
        p.add_run(texto)


def add_filas(tabla, n):
    """Clona la última fila n veces, limpiando su contenido."""
    for _ in range(n):
        nueva = copy.deepcopy(tabla.rows[-1]._tr)
        tabla.rows[-1]._tr.addnext(nueva)
        for c in tabla.rows[-1].cells:
            _set(c, "")


def conclusiones(doc, textos):
    """Localiza los párrafos '1.-','2.-','3.-' posteriores a CONCLUSIONES y los llena."""
    inicio = None
    for i, p in enumerate(doc.paragraphs):
        if "CONCLUSI" in p.text.upper():
            inicio = i
            break
    if inicio is None:
        raise RuntimeError("no encontré CONCLUSIONES")
    objetivo, n = [], 1
    for i in range(inicio, len(doc.paragraphs)):
        t = doc.paragraphs[i].text.strip()
        if t.startswith(f"{n}.-") or t == f"{n}." or t.startswith(f"{n}. "):
            objetivo.append(i); n += 1
            if n > len(textos): break
    if len(objetivo) != len(textos):
        raise RuntimeError(f"esperaba {len(textos)} marcadores, encontré {len(objetivo)}")
    for idx, txt in zip(objetivo, textos):
        set_parrafo(doc, idx, txt)
    return objetivo
